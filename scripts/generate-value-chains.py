"""Generate value-chain registry and unique sector schemas from the field dictionary."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "scripts" / "field-dictionary.xlsx"
if not XLSX.exists():
    XLSX = Path(r"c:\Users\FRANC\Downloads\MkulimaCollect_Field_Dictionary_and_Value_Chain_Map.xlsx")
OUT_REG = ROOT / "apps" / "collect" / "features" / "sectors" / "valueChainRegistry.ts"
OUT_GEN = ROOT / "apps" / "collect" / "features" / "sectors" / "generatedChainSchemas.ts"
OUT_IDS = ROOT / "apps" / "collect" / "constants" / "sectorIds.ts"
OUT_WEB = ROOT / "src" / "generatedWebSectors.ts"

HANDMADE = {
    "VC01": "dairy",
    "VC04": "poultry",
    "VC06": "aquaculture",
    "VC08": "maize",
    "VC09": "beans",
    "VC10": "rice",
    "VC11": "irish-potato",
    "VC25": "tea",
    "VC26": "coffee",
    "VC27": "avocado",
    "VC28": "macadamia",
    "VC34": "tomato",
}

ALIASES = {
    "livestock-meat": "beef-cattle",
    "horticulture": "horticulture",
}

GROUP_MAP = {
    "Livestock": "Livestock",
    "Aquaculture": "Aquaculture",
    "Apiculture": "Livestock",
    "Annual crop": "Annual crops",
    "Perennial crop": "Perennial crops",
    "Horticulture": "Horticulture",
}

SCORE_MAP = {
    "Livestock": "Livestock production score",
    "Aquaculture": "Aquaculture production score",
    "Apiculture": "Livestock production score",
    "Annual crop": "Crop production score",
    "Perennial crop": "Perennial crop score",
    "Horticulture": "Horticulture score",
}

EVIDENCE = {
    "Livestock": [
        ("animal-photo", "Animal photograph"),
        ("sales-record", "Sales record"),
        ("veterinary-record", "Veterinary record"),
        ("movement-permit", "Movement permit"),
    ],
    "Aquaculture": [
        ("pond-photo", "Pond photograph"),
        ("feed-receipt", "Feed receipt"),
        ("harvest-record", "Harvest record"),
        ("sales-record", "Sales record"),
    ],
    "Annual crops": [
        ("farm-photo", "Farm photograph"),
        ("input-receipt", "Input receipt"),
        ("harvest-record", "Harvest record"),
        ("buyer-delivery-note", "Buyer delivery note"),
    ],
    "Perennial crops": [
        ("farm-photo", "Farm photograph"),
        ("factory-delivery-slip", "Factory delivery slip"),
        ("payment-statement", "Payment statement"),
        ("input-receipt", "Input receipt"),
    ],
    "Horticulture": [
        ("farm-photo", "Farm photograph"),
        ("spray-record", "Spray record"),
        ("harvest-record", "Harvest record"),
        ("buyer-delivery-note", "Buyer delivery note"),
    ],
}


def slugify(code: str, name: str) -> str:
    if code in HANDMADE:
        return HANDMADE[code]
    return (
        name.lower()
        .replace(" and ", "-")
        .replace("/", "-")
        .replace("&", "and")
        .replace(",", "")
        .replace("(", "")
        .replace(")", "")
        .replace(" ", "-")
    )


def camel(value: str) -> str:
    parts = [part for part in re.split(r"[^a-zA-Z0-9]+", value) if part]
    if not parts:
        return "field"
    head, *tail = parts
    return head[:1].lower() + head[1:] + "".join(part[:1].upper() + part[1:] for part in tail)


def const_key(slug: str) -> str:
    return slug.replace("-", "_").upper()


def ts_str(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def parse_options(data_type: str | None, unit_or_options: str | None, validation: str | None) -> list[str] | None:
    raw = " ".join(part for part in (unit_or_options, validation) if part)
    if "|" not in raw:
        return None
    opts = []
    for part in raw.split("|"):
        item = part.strip()
        if not item or len(item) > 36:
            continue
        if " " in item and len(item.split()) > 4:
            continue
        opts.append(item.replace("_", " "))
    if 2 <= len(opts) <= 16:
        return opts
    return None


def parse_unit(data_type: str | None, unit_or_options: str | None) -> str | None:
    value = (unit_or_options or "").strip()
    if not value or "|" in value or len(value) > 18:
        return None
    if (data_type or "").lower() in {"enum", "enum/string", "multi_enum", "boolean"}:
        return None
    return value


def map_type(data_type: str | None, options: list[str] | None) -> tuple[str, bool]:
    raw = (data_type or "text").lower()
    if raw in {"repeat", "object"}:
        return "text", True
    if raw in {"integer"}:
        return "integer", False
    if raw in {"decimal"}:
        return "decimal", False
    if raw in {"money"}:
        return "currency", False
    if raw in {"date", "datetime", "year", "date_or_year"}:
        return "date", False
    if raw in {"boolean"}:
        return "yes_no", False
    if raw in {"enum", "enum/string"}:
        return ("single_choice" if options else "text"), False
    if raw in {"multi_enum"}:
        return ("multiple_choice" if options else "text"), False
    return "text", False


def load_workbook():
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    chains = []
    ws = wb["Value Chains"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1 or not row or not row[0] or not str(row[0]).startswith("VC"):
            continue
        code, name, sector, subtypes, status, cadence = (row + (None,) * 6)[:6]
        chains.append(
            {
                "code": str(code),
                "label": str(name),
                "sector": str(sector),
                "group": GROUP_MAP.get(str(sector), "Annual crops"),
                "subtypes": str(subtypes or ""),
                "status": str(status or "standard"),
                "cadence": str(cadence or ""),
                "slug": slugify(str(code), str(name)),
                "scorePath": SCORE_MAP.get(str(sector), "General agriculture score"),
            }
        )

    dict_ws = wb["Field Dictionary"]
    header = None
    by_scope: dict[str, list[dict]] = defaultdict(list)
    for i, row in enumerate(dict_ws.iter_rows(values_only=True), 1):
        vals = [None if v is None else str(v).strip() for v in row]
        if i == 4:
            header = vals
            continue
        if not header or not any(vals):
            continue
        rec = dict(zip(header, vals + [None] * (len(header) - len(vals))))
        scope = rec.get("Scope") or ""
        if scope:
            by_scope[scope].append(rec)
    wb.close()
    return chains, by_scope


def field_objects(chain: dict, records: list[dict]) -> dict[str, list[dict]]:
    prefix = camel(chain["slug"])
    fields_by_module: dict[str, list[dict]] = defaultdict(list)
    seen = set()
    for rec in records:
        field_id = rec.get("Field ID") or ""
        if not field_id or field_id in seen:
            continue
        seen.add(field_id)
        unique_id = prefix + camel(field_id)[:1].upper() + camel(field_id)[1:]
        label = rec.get("Label") or field_id.replace("_", " ").title()
        module = rec.get("Module") or "Production"
        data_type = rec.get("Data type")
        options = parse_options(data_type, rec.get("Unit or options"), rec.get("Validation"))
        unit = parse_unit(data_type, rec.get("Unit or options"))
        required = (rec.get("Required when") or "").lower() == "always"
        mapped, repeatable = map_type(data_type, options)
        field = {
            "id": unique_id,
            "label": label,
            "type": mapped,
            "required": required,
        }
        if unit:
            field["unit"] = unit
        if options:
            field["options"] = options
        if repeatable:
            field["repeatable"] = True
        fields_by_module[module].append(field)
    return fields_by_module


def emit_field(field: dict) -> str:
    parts = [f"id: {ts_str(field['id'])}", f"label: {ts_str(field['label'])}", f"type: {ts_str(field['type'])}"]
    if field.get("required"):
        parts.append("required: true")
    if field.get("repeatable"):
        parts.append("repeatable: true")
    if field.get("unit"):
        parts.append(f"unit: {ts_str(field['unit'])}")
    if field.get("options"):
        parts.append("options: [" + ", ".join(ts_str(option) for option in field["options"]) + "]")
    return "{ " + ", ".join(parts) + " }"


def evidence_fields(chain: dict) -> list[dict]:
    prefix = camel(chain["slug"])
    if chain["slug"] == "apiculture":
        items = [
            ("hive-photo", "Hive photograph"),
            ("harvest-record", "Harvest record"),
            ("sales-record", "Sales record"),
            ("inspection-record", "Inspection record"),
        ]
    else:
        items = EVIDENCE.get(chain["group"], EVIDENCE["Annual crops"])
    return [
        {
            "id": prefix + camel(item_id)[:1].upper() + camel(item_id)[1:],
            "label": f"{chain['label']} {label[0].lower() + label[1:]}",
            "type": "evidence",
            "required": True,
        }
        for item_id, label in items
    ]


def main() -> None:
    chains, by_scope = load_workbook()

    lines = [
        "export type ValueChainStatus = \"priority\" | \"standard\";",
        "export type ValueChainGroup = \"Annual crops\" | \"Horticulture\" | \"Perennial crops\" | \"Livestock\" | \"Aquaculture\";",
        "",
        "export type ValueChain = {",
        "  code: string;",
        "  slug: string;",
        "  label: string;",
        "  group: ValueChainGroup;",
        "  sector: string;",
        "  subtypes: string;",
        "  status: ValueChainStatus;",
        "  cadence: string;",
        "  scorePath: string;",
        "};",
        "",
        "export const valueChains: readonly ValueChain[] = [",
    ]
    for chain in chains:
        lines.append(
            "  { "
            + f"code: {ts_str(chain['code'])}, slug: {ts_str(chain['slug'])}, label: {ts_str(chain['label'])}, "
            + f"group: {ts_str(chain['group'])}, sector: {ts_str(chain['sector'])}, subtypes: {ts_str(chain['subtypes'])}, "
            + f"status: {ts_str(chain['status'])}, cadence: {ts_str(chain['cadence'])}, scorePath: {ts_str(chain['scorePath'])}"
            + " },"
        )
    lines.append("];")
    lines.append("")
    lines.append("export const handmadeSectorSlugs = new Set([")
    for slug in sorted(set(HANDMADE.values()) | set(ALIASES)):
        lines.append(f"  {ts_str(slug)},")
    lines.append("]);")
    lines.append("")
    lines.append("export const sectorAliases: Record<string, string> = {")
    for alias, target in ALIASES.items():
        lines.append(f"  {ts_str(alias)}: {ts_str(target)},")
    lines.append("};")
    lines.append("")
    lines.append("export function getValueChainBySlug(slug: string): ValueChain | undefined {")
    lines.append("  const resolved = sectorAliases[slug] ?? slug;")
    lines.append("  return valueChains.find((chain) => chain.slug === resolved || chain.slug === slug);")
    lines.append("}")
    lines.append("")
    OUT_REG.write_text("\n".join(lines) + "\n", encoding="utf-8")

    slugs = [chain["slug"] for chain in chains] + list(ALIASES)
    unique_slugs = list(dict.fromkeys(slugs))
    id_lines = ["export const SectorId = {"]
    for slug in unique_slugs:
        id_lines.append(f"  {const_key(slug)}: {ts_str(slug)},")
    id_lines.append("} as const;")
    id_lines.append("")
    id_lines.append("export type SectorIdValue = (typeof SectorId)[keyof typeof SectorId];")
    id_lines.append("")
    OUT_IDS.write_text("\n".join(id_lines) + "\n", encoding="utf-8")

    gen = [
        "import type { SectorSchemaDefinition } from \"./types\";",
        "",
        "function schema(sector: string, title: string, sections: SectorSchemaDefinition[\"sections\"]): SectorSchemaDefinition {",
        "  return { id: `${sector}-field-v1`, version: \"1.0.0\", sector, title: `${title} collection`, sections };",
        "}",
        "",
        "export const generatedChainSchemas: Record<string, SectorSchemaDefinition> = {",
    ]
    web = [
        "export const generatedWebSectors = [",
    ]

    generated_count = 0
    for chain in chains:
        if chain["code"] in HANDMADE:
            continue
        fields_by_module = field_objects(chain, by_scope.get(chain["code"], []))
        if not fields_by_module:
            continue
        generated_count += 1
        sections_ts = []
        web_sections = []
        for module, fields in fields_by_module.items():
            section_id = re.sub(r"[^a-z0-9]+", "-", module.lower()).strip("-")[:48]
            field_objs = [emit_field(field) for field in fields]
            sections_ts.append(
                "{ id: "
                + ts_str(section_id)
                + ", title: "
                + ts_str(module)
                + ", fields: [\n        "
                + ",\n        ".join(field_objs)
                + "\n      ] }"
            )
            web_fields = []
            for field in fields:
                web_type = "number" if field["type"] in {"integer", "decimal", "currency", "quantity"} else "choice" if field["type"] in {"single_choice", "multiple_choice", "yes_no"} else "text"
                parts = [f"id: {ts_str(field['id'])}", f"label: {ts_str(field['label'])}"]
                if web_type != "text":
                    parts.append(f"type: {ts_str(web_type)}")
                if field.get("required"):
                    parts.append("required: true")
                if field.get("options"):
                    parts.append("options: [" + ", ".join(ts_str(option) for option in field["options"]) + "]")
                elif field["type"] == "yes_no":
                    parts.append('options: ["Yes", "No"]')
                web_fields.append("{ " + ", ".join(parts) + " }")
            web_sections.append("{ title: " + ts_str(module) + ", fields: [\n        " + ",\n        ".join(web_fields) + "\n      ] }")

        evidence = evidence_fields(chain)
        sections_ts.append(
            "{ id: \"evidence\", title: \"Required evidence\", fields: [\n        "
            + ",\n        ".join(emit_field(field) for field in evidence)
            + "\n      ] }"
        )
        gen.append(
            f"  {ts_str(chain['slug'])}: schema({ts_str(chain['slug'])}, {ts_str(chain['label'])}, [\n      "
            + ",\n      ".join(sections_ts)
            + "\n    ]),"
        )
        web.append(
            "  {\n    id: "
            + ts_str(chain["slug"])
            + ",\n    label: "
            + ts_str(chain["label"])
            + ",\n    group: "
            + ts_str(chain["group"])
            + ",\n    status: "
            + ts_str(chain["status"])
            + ",\n    evidence: ["
            + ", ".join(ts_str(field["label"]) for field in evidence)
            + "],\n    sections: [\n      "
            + ",\n      ".join(web_sections)
            + "\n    ],\n  },"
        )

    gen.append("};")
    gen.append("")
    web.append("];")
    web.append("")
    OUT_GEN.write_text("\n".join(gen) + "\n", encoding="utf-8")
    OUT_WEB.write_text("\n".join(web) + "\n", encoding="utf-8")
    print("chains", len(chains), "generated", generated_count)
    print("wrote", OUT_REG)
    print("wrote", OUT_GEN)
    print("wrote", OUT_IDS)
    print("wrote", OUT_WEB)


if __name__ == "__main__":
    main()
